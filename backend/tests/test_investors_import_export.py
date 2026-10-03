"""Backend tests for IPO Premium accounts import/export (LotNeeti)."""
import io
import os
import uuid
import requests
import pytest
from openpyxl import load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    with open("/app/frontend/.env") as f:
        for line in f:
            if line.startswith("REACT_APP_BACKEND_URL="):
                BASE_URL = line.split("=", 1)[1].strip().rstrip("/")

EXPECTED_COLUMNS = ["Name", "PAN", "Type", "DPID", "CLIENT ID", "UPI ID", "Account Number", "Bank Name"]


@pytest.fixture(scope="module")
def fresh_user_token():
    email = f"test_import_{uuid.uuid4().hex[:8]}@test.com"
    r = requests.post(f"{BASE_URL}/api/auth/register",
                      json={"name": "Imp Tester", "email": email, "password": "pw123456"})
    assert r.status_code == 200, r.text
    return r.json()["token"], email


@pytest.fixture
def auth_headers(fresh_user_token):
    token, _ = fresh_user_token
    return {"Authorization": f"Bearer {token}"}


CSV_SAMPLE = (
    "Name,PAN,Type,DPID,CLIENT ID,UPI ID,Account Number,Bank Name\n"
    "RAM GOPAL PATIDAR HUF,ABBHR1134L,CDSL,,1208870432110422,gal.pat1968@ybl,,\n"
    "SITA DEVI,ABCDE1234F,NSDL,IN300123,00001111,sita@upi,123456789,HDFC\n"
    "GOPAL,FGHIJ5678K,CDSL,,22223333,gopal@upi,987654321,ICICI\n"
    # in-file dup PAN (same as row 2)
    "DUPLICATE PAN,ABCDE1234F,NSDL,,99999999,dup@upi,,\n"
    # invalid: missing PAN
    "NO PAN GUY,,CDSL,,1234,none@upi,,\n"
)


def test_import_mixed_rows(auth_headers):
    files = {"file": ("accounts.csv", CSV_SAMPLE.encode("utf-8"), "text/csv")}
    r = requests.post(f"{BASE_URL}/api/investors/import", headers=auth_headers, files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data == {"created": 3, "skipped_duplicates": 1, "invalid": 1, "total_rows": 5}, data


def test_reimport_same_file_all_skipped(auth_headers):
    files = {"file": ("accounts.csv", CSV_SAMPLE.encode("utf-8"), "text/csv")}
    r = requests.post(f"{BASE_URL}/api/investors/import", headers=auth_headers, files=files)
    assert r.status_code == 200
    data = r.json()
    # 4 valid PAN rows all exist now; 1 invalid still invalid
    assert data["created"] == 0
    assert data["skipped_duplicates"] == 4
    assert data["invalid"] == 1


def test_investors_listed_after_import(auth_headers):
    r = requests.get(f"{BASE_URL}/api/investors", headers=auth_headers)
    assert r.status_code == 200
    pans = {i["pan"] for i in r.json()}
    assert {"ABBHR1134L", "ABCDE1234F", "FGHIJ5678K"}.issubset(pans)


def test_export_roundtrip(auth_headers):
    r = requests.get(f"{BASE_URL}/api/investors/export", headers=auth_headers)
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    assert header == EXPECTED_COLUMNS, f"Header mismatch: {header}"
    data_rows = rows[1:]
    # Unique PANs only
    pans = [row[1] for row in data_rows]
    assert len(pans) == len(set(pans)), f"Duplicate PANs in export: {pans}"
    # Our 3 imported PANs should be in export
    assert {"ABBHR1134L", "ABCDE1234F", "FGHIJ5678K"}.issubset(set(pans))
    # UPI ID preserved (col idx 5)
    by_pan = {row[1]: row for row in data_rows}
    assert by_pan["ABBHR1134L"][5] == "gal.pat1968@ybl"
    assert by_pan["ABCDE1234F"][5] == "sita@upi"
    # Type preserved
    assert by_pan["ABCDE1234F"][2] == "NSDL"


def test_import_rejects_unknown_format(auth_headers):
    files = {"file": ("x.txt", b"hello", "text/plain")}
    r = requests.post(f"{BASE_URL}/api/investors/import", headers=auth_headers, files=files)
    assert r.status_code == 400


def test_import_requires_pan_column(auth_headers):
    bad = "Name,Type\nFoo,CDSL\n"
    files = {"file": ("bad.csv", bad.encode(), "text/csv")}
    r = requests.post(f"{BASE_URL}/api/investors/import", headers=auth_headers, files=files)
    assert r.status_code == 400


def test_export_requires_auth():
    r = requests.get(f"{BASE_URL}/api/investors/export")
    assert r.status_code == 401
